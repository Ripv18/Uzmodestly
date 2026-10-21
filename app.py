
import os
import io
import json
import time
import base64
import logging
import uuid
from collections import defaultdict, deque

import requests
from flask import Flask, send_from_directory, request, jsonify, Response, stream_with_context
from werkzeug.utils import secure_filename

import memory

try:
    from PIL import Image, ImageOps
except ImportError:
    Image = None
    ImageOps = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None


app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 12 * 1024 * 1024
app.json.ensure_ascii = False

logging.basicConfig(level=logging.INFO)

GROQ_URL = "https://api.groq.com/openai/v1/chat/completions"
MODEL = os.environ.get("GROQ_MODEL", "openai/gpt-oss-20b")
VISION_MODEL = os.environ.get("GROQ_VISION_MODEL", "qwen/qwen3.8-27b")

IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp"}
DOCUMENT_EXTENSIONS = {"pdf", "docx"}
requests_by_ip = defaultdict(deque)


def get_user_id():
    raw_id = request.headers.get("X-User-ID", "").strip()
    try:
        return str(uuid.UUID(raw_id))
    except (ValueError, AttributeError, TypeError):
        return None


def repair_text(text):
    if not isinstance(text, str):
        return text

    replacements = {
        "â\x80\x98": "‘",
        "â\x80\x99": "’",
        "â": "‘",
        "â": "’",
        "â€“": "–",
        "â€”": "—",
        "â€¦": "…",
    }

    for wrong, correct in replacements.items():
        text = text.replace(wrong, correct)

    return text


def allowed_file(filename):
    if "." not in filename:
        return False

    ext = filename.rsplit(".", 1)[1].lower()
    return ext in IMAGE_EXTENSIONS | DOCUMENT_EXTENSIONS


def extract_document(filename, raw_bytes):
    ext = filename.rsplit(".", 1)[1].lower()

    if ext == "pdf":
        if PdfReader is None:
            raise ValueError("PDF uchun pypdf kutubxonasi kerak.")

        reader = PdfReader(io.BytesIO(raw_bytes))

        if reader.is_encrypted:
            raise ValueError("Parolli PDF hozircha qo‘llanmaydi.")

        parts = [
            page.extract_text() or ""
            for page in reader.pages[:50]
        ]

        text = "\n".join(parts).strip()

        if not text:
            raise ValueError(
                "PDF ichidan matn topilmadi. "
                "Skanerlangan hujjatni rasm sifatida yubor."
            )

        return text[:40000]

    if ext == "docx":
        if Document is None:
            raise ValueError("Word fayli uchun python-docx kerak.")

        document = Document(io.BytesIO(raw_bytes))
        parts = [
            p.text for p in document.paragraphs if p.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        text = "\n".join(parts).strip()

        if not text:
            raise ValueError("Word faylidan matn topilmadi.")

        return text[:40000]

    raise ValueError("Bu hujjat turi qo‘llab-quvvatlanmaydi.")


def read_attachment(file):
    filename = secure_filename(file.filename or "")

    if not filename or not allowed_file(filename):
        raise ValueError(
            "Faqat PNG, JPG, JPEG, WEBP, PDF yoki DOCX yubor."
        )

    raw_bytes = file.read()

    if not raw_bytes:
        raise ValueError("Tanlangan fayl bo‘sh.")

    if len(raw_bytes) > 10 * 1024 * 1024:
        raise ValueError("Fayl hajmi 10 MB dan oshmasligi kerak.")

    ext = filename.rsplit(".", 1)[1].lower()

    if ext in IMAGE_EXTENSIONS:
        if Image is None or ImageOps is None:
            raise ValueError("Rasm uchun Pillow kutubxonasi kerak.")

        try:
            with Image.open(io.BytesIO(raw_bytes)) as image:
                image.verify()

            with Image.open(io.BytesIO(raw_bytes)) as image:
                image = ImageOps.exif_transpose(image)
                image = image.convert("RGB")
                image.thumbnail((3072, 3072))

                buffer = io.BytesIO()
                image.save(
                    buffer,
                    format="JPEG",
                    quality=92,
                    optimize=True,
                )

            encoded = base64.b64encode(
                buffer.getvalue()
            ).decode("ascii")

        except Exception as exc:
            raise ValueError("Rasmni o‘qib bo‘lmadi.") from exc

        return {
            "name": filename,
            "type": "image",
            "data_url": "data:image/jpeg;base64," + encoded,
        }

    return {
        "name": filename,
        "type": "document",
        "text": extract_document(filename, raw_bytes),
    }


@app.errorhandler(413)
def file_too_large(_error):
    return jsonify({
        "error": "Fayl juda katta. Kichikroq fayl yubor."
    }), 413


@app.route("/")
def home():
    return send_from_directory(".", "index.html")


@app.route("/google60ac72fbd987c0b6.html")
def google_verify():
    return send_from_directory(".", "google60ac72fbd987c0b6.html")


@app.route("/api/conversations", methods=["GET"])
def conversations():
    user_id = get_user_id()
user_name = request.headers.get("X-User-Name", "").strip()[:50]
    if not user_id:
        return jsonify({
            "error": "Foydalanuvchi ID noto‘g‘ri yoki yo‘q."
        }), 400

    return jsonify(
        memory.list_conversations(user_id=user_id)
    )


@app.route("/api/conversations", methods=["POST"])
def new_conversation():
    user_id = get_user_id()

    if not user_id:
        return jsonify({
            "error": "Foydalanuvchi ID noto‘g‘ri yoki yo‘q."
        }), 400

    try:
        cid = memory.create_conversation(user_id=user_id)

        return jsonify({
            "id": cid,
            "title": "Yangi suhbat",
        })

    except Exception:
        app.logger.exception("Suhbat yaratishda xato")

        return jsonify({
            "error": "Suhbat yaratilmadi."
        }), 500


@app.route("/api/conversations/<int:cid>", methods=["GET"])
def conversation_detail(cid):
    user_id = get_user_id()

    if not user_id:
        return jsonify({
            "error": "Foydalanuvchi ID noto‘g‘ri yoki yo‘q."
        }), 400

    if not memory.conversation_belongs_to(cid, user_id):
        return jsonify({
            "error": "Suhbat topilmadi."
        }), 404

    return jsonify({
        "id": cid,
        "messages": memory.get_messages(
            cid,
            limit=1000,
            user_id=user_id,
        ),
    })


@app.route("/api/conversations/<int:cid>", methods=["DELETE"])
def delete_conversation(cid):
    user_id = get_user_id()

    if not user_id:
        return jsonify({
            "error": "Foydalanuvchi ID noto‘g‘ri yoki yo‘q."
        }), 400

    try:
        deleted = memory.delete_conversation(
            cid,
            user_id=user_id,
        )

        if not deleted:
            return jsonify({
                "error": "Suhbat topilmadi."
            }), 404

        return jsonify({
            "success": True,
            "message": "Suhbat o‘chirildi.",
        })

    except Exception:
        app.logger.exception("Suhbatni o‘chirishda xato")

        return jsonify({
            "error": "Suhbat o‘chirilmadi."
        }), 500


@app.route("/api/chat", methods=["POST"])
def chat():
    user_id = get_user_id()
    user_name = request.headers.get("X-User-Name", "").strip()[:50]

    if not user_id:
        return jsonify({
            "error": "Foydalanuvchi ID noto‘g‘ri yoki yo‘q."
        }), 400

    api_key = os.environ.get("GROQ_API_KEY", "").strip()

    if not api_key:
        return jsonify({
            "error": "GROQ_API_KEY sozlanmagan."
        }), 503
    question = (
        request.form.get("message") or ""
    ).strip()[:3000]

    uploaded_file = request.files.get("file")
    has_file = bool(
        uploaded_file and uploaded_file.filename
    )

    if not question and not has_file:
        return jsonify({
            "error": "Xabar yoz yoki fayl yubor."
        }), 400

    raw_cid = request.form.get("conversation_id")

    try:
        cid = int(raw_cid) if raw_cid else None
    except (TypeError, ValueError):
        return jsonify({
            "error": "Suhbat ID noto‘g‘ri."
        }), 400

    attachment = None

    if has_file:
        try:
            attachment = read_attachment(uploaded_file)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400

    if cid is None:
        try:
            cid = memory.create_conversation(
                user_id=user_id
            )
        except Exception:
            app.logger.exception("Suhbat yaratilmadi")

            return jsonify({
                "error": "Suhbat yaratilmadi."
            }), 500

    if not memory.conversation_belongs_to(cid, user_id):
        return jsonify({
            "error": "Suhbat topilmadi."
        }), 404

    ip = request.remote_addr or "unknown"
    now = time.time()
    recent = requests_by_ip[ip]

    while recent and now - recent[0] > 60:
        recent.popleft()

    if len(recent) >= 10:
        return jsonify({
            "error": "Limit tugadi. Bir daqiqadan keyin urin."
        }), 429

    system_prompt = """
SENING SHAXSING VA JAVOB BERISH QOIDALARI:

1. Sening noming Uzmodestly AI.
2. "Isming nima?" deyilsa, "Men Uzmodestly AI man" deb javob ber.
3. "Seni kim yaratgan?" yoki "Xo'jayining kim?" deyilsa,
   "Meni Behruz Tojimurodov yaratgan" deb javob ber.
4. O'zingni ChatGPT deb tanishtirma.
5. O'zbek tilida tabiiy, sodda va tushunarli javob ber.
6. Oddiy va xavfsiz savollarga asossiz ravishda
   "I can't help with that" deb javob berma.
7. Javob bera olmasang, sababini o'zbek tilida tushuntir.
8. Matematik masalalarni tekshirib, bosqichma-bosqich yech.
9. Rasm yuborilsa, undagi yozuvlarni diqqat bilan tahlil qil.
10. Rasm xira bo'lsa, buni ochiq ayt.
11. Bilmagan narsangni to'qib chiqarma.
12. Oddiy javoblarda ortiqcha bezaklardan foydalanma.
13. Kod so'ralsa, to'liq va ishlaydigan kod ber.
"""

    if user_name:
        system_prompt += (
            "\n\nFOYDALANUVCHI HAQIDA:"
            f"\nFoydalanuvchining ismi: {user_name}."
            "\nIsmi so'ralsa, aynan shu ismni ayt."
            "\nIsmini oldin aytmagan deb aytma."
        )

    safe_messages = [
        {"role": "system", "content": system_prompt}
    ]
    try:
        old_context = memory.get_recent_context(
            exclude_conversation_id=cid,
            limit=12,
            user_id=user_id,
        )

        if old_context:
            context_text = "\n".join(
                f"[{item['title']}] "
                f"{item['role']}: {item['content']}"
                for item in old_context
            )

            safe_messages.append({
                "role": "system",
                "content": (
                    "Ushbu foydalanuvchining oldingi "
                    "suhbatlaridan olingan kontekst. "
                    "Tegishli bo‘lsa foydalan:\n"
                    + context_text[:6000]
                ),
            })

        for item in memory.get_messages(
            cid,
            limit=30,
            user_id=user_id,
        ):
            if item["role"] in ("user", "assistant"):
                safe_messages.append({
                    "role": item["role"],
                    "content": item["content"][:3000],
                })

    except Exception:
        app.logger.exception("Xotirani o‘qishda xato")

        return jsonify({
            "error": "Suhbat xotirasini o‘qib bo‘lmadi."
        }), 500

    saved_question = question or "Yuborilgan faylni tahlil qil."
    user_content = saved_question

    if attachment:
        saved_question = (
            f"[Fayl: {attachment['name']}]\n"
            f"{saved_question}"
        )

        if attachment["type"] == "document":
            user_content += (
                f"\n\nFayl: {attachment['name']}\n"
                "Fayldan ajratilgan matn:\n"
                + attachment["text"]
            )

        else:
            user_content = [
                {
                    "type": "text",
                    "text": (
                        user_content
                        + "\n\nRASM TAHLILI QOIDALARI:"
                        + "\n1. Rasm matnini diqqat bilan o‘qi."
                        + "\n2. Sonlar, ishoralar, kasrlar va formulalarni tekshir."
                        + "\n3. Masala shartini o‘zgartirma."
                        + "\n4. Noaniq belgilarni taxmin qilma."
                        + "\n5. Masalani bosqichma-bosqich yech."
                        + "\n6. Javobni qayta tekshir."
                    ),
                },
                {
                    "type": "image_url",
                    "image_url": {
                        "url": attachment["data_url"]
                    },
                },
            ]

    try:
        memory.save_message(
            cid,
            "user",
            saved_question,
            user_id=user_id,
        )

        safe_messages.append({
            "role": "user",
            "content": user_content,
        })

    except Exception:
        app.logger.exception("Xabarni saqlashda xato")

        return jsonify({
            "error": "Xabar saqlanmadi."
        }), 500

    recent.append(now)

    selected_model = (
        VISION_MODEL
        if attachment and attachment["type"] == "image"
        else MODEL
    )

    def generate():
        reply_parts = []

        try:
            with requests.post(
                GROQ_URL,
                headers={
                    "Authorization": f"Bearer {api_key}",
                    "Content-Type": "application/json",
                },
                json={
                    "model": selected_model,
                    "messages": safe_messages[-31:],
                    "temperature": 0.4,
                    "max_tokens": 1200,
                    "stream": True,
                },
                stream=True,
                timeout=(10, 120),
            ) as response:

                if not response.ok:
                    app.logger.error(
                        "Groq API HTTP %s: %s",
                        response.status_code,
                        response.text[:500],
                    )

                    yield (
                        "AI API xatosi. Model nomi, limit yoki "
                        "API kalit sozlamalarini tekshir."
                    )
                    return

                for line in response.iter_lines(
                    decode_unicode=False
                ):
                    if not line:
                        continue

                    line = line.decode(
                        "utf-8",
                        errors="replace",
                    )

                    if not line.startswith("data:"):
                        continue

                    payload = line[5:].strip()

                    if payload == "[DONE]":
                        break

                    try:
                        obj = json.loads(payload)
                        part = (
                            obj["choices"][0]["delta"]
                            .get("content")
                        )

                        if isinstance(part, str) and part:
                            reply_parts.append(part)

                    except (
                        ValueError,
                        KeyError,
                        IndexError,
                        TypeError,
                    ):
                        continue

            reply = repair_text(
                "".join(reply_parts)
            ).strip()

            if not reply:
                yield "Javob olinmadi. Qayta urinib ko‘r."
                return

            memory.save_message(
                cid,
                "assistant",
                reply,
                user_id=user_id,
            )

            yield reply

        except requests.RequestException:
            app.logger.exception("Groq API ulanish xatosi")
            yield (
                "AI serveriga ulanib bo‘lmadi. "
                "Keyinroq urin."
            )

        except Exception:
            app.logger.exception("Javob tayyorlashda xato")
            yield "Javob tayyorlashda xato yuz berdi."

    return Response(
        stream_with_context(generate()),
        content_type="text/plain; charset=utf-8",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
        },
    )


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    app.run(
        host="0.0.0.0",
        port=port,
        debug=False,
    )
